# -*- coding: utf-8 -*-
"""Add the enclosure fit-check slides (off-board parts inside the case) to both v1.4 decks."""
import importlib.util, os, sys
spec = importlib.util.spec_from_file_location("ud", os.path.join(os.path.dirname(__file__), "update_decks_v14.py"))
ud = importlib.util.module_from_spec(spec); spec.loader.exec_module(ud)
from pptx import Presentation
E, tbox, pic, rect, square, title = ud.E, ud.tbox, ud.pic, ud.rect, ud.square, ud.title
INK, INK2, TEAL, PINK, AMBER = ud.INK, ud.INK2, ud.TEAL, ud.PINK, ud.AMBER

# ---- full deck: new slide 16 after "外殼各面，與怎麼印"
prs = Presentation("esp32s3-camera-v1.4.pptx")
assert len(prs.slides) == 19
s = ud.add_slide_at(prs, 15)
title(s, "外殼裝配檢查：板外零件放進去看", "相機模組與反折排線、電池與電池線、記憶卡、螢幕模組都依文件尺寸建成簡化模型放進 SolidWorks 組合件，和兩片殼一起跑干涉檢查：0 件。殼拿掉或調半透明拍的。")
views = [("fit-front-open-iso.png", "前殼拿掉", "黑色相機模組與從 J2 反折回來貼在板面的排線；反折處離上牆內面 3.8 mm，鏡座頂離天花板 1.1 mm。上板邊的快門開關對著上牆的窗。"),
         ("fit-back-open.png", "後殼拿掉", "藍色螢幕模組插在 J5 母座上、蓋在銀色電池上方；電池線從尾端繞到 J4 插頭；記憶卡從左板邊露出，正好穿過後殼的卡槽。"),
         ("fit-ghost-back.png", "兩片殼半透明", "後殼四支定位柱穿進螢幕四角的 Ø2 孔、頂在模組 PCB 朝底板那一面；四支立柱頂住板子四角；螺絲從背面沉孔鎖穿。")]
for i, (f, hd, bd) in enumerate(views):
    x = 0.60 + i * 4.10
    pic(s, E(f), x, 1.70, 3.90)
    tbox(s, x, 4.20, 3.90, 0.32, hd, 14, True, INK)
    tbox(s, x, 4.55, 3.90, 1.20, bd, 10.5, False, INK2, space_after=0)
rect(s, 0.60, 5.85, 12.10, 0.95)
square(s, 0.85, 6.05, TEAL, 0.20)
tbox(s, 1.15, 6.00, 11.30, 0.30, "模型裡刻意沒放的東西", 12, True, INK)
tbox(s, 1.15, 6.32, 11.30, 0.45, "電池底下的泡棉（會包住 1.3 mm 高的零件，放進去只會報假警報）、插在 J4 裡的插頭本體、插在 J3 裡的卡身、母座裡的針腳。相機鏡座原本只離天花板 0.1 mm，所以前殼內高從 7 改成 8 mm。", 9.5, False, INK2, space_after=0)
ud.finish(prs, ud.FOOTER_FULL, page_x=11.70)
prs.save("esp32s3-camera-v1.4.pptx"); print("full deck:", len(prs.slides), "slides")

# ---- simple deck: new slide after "外殼長這樣"
prs = Presentation("esp32s3-camera-v1.4-簡易版.pptx")
assert len(prs.slides) == 10
s = ud.add_slide_at(prs, 9)
tbox(s, 0.70, 0.50, 12.0, 0.80, "裝進殼裡長這樣", 30, True, INK)
tbox(s, 0.70, 1.30, 12.0, 0.50, "把鏡頭、排線、電池、記憶卡、螢幕都放進電腦模型裡試裝過，什麼都沒有卡到。這兩張是把殼拿掉一半看裡面。", 13, False, INK2)
pic(s, E("fit-front-open-iso.png"), 0.70, 1.95, 5.90)
pic(s, E("fit-back-open.png"), 6.80, 1.95, 5.90)
tbox(s, 0.70, 5.70, 5.90, 0.30, "正面那片殼拿掉：鏡頭和它的排線貼在板子上", 11, False, INK2)
tbox(s, 6.80, 5.70, 5.90, 0.30, "背面那片殼拿掉：螢幕插在插座上、蓋在電池上面", 11, False, INK2)
for i, (hd, bd, col) in enumerate([("鏡頭", "排線從板子上緣的座子出來折回來，鏡頭貼在板面上，離殼頂還有一點空間。", TEAL),
                                   ("螢幕和電池", "螢幕的針直接插進板子右邊的插座，電池夾在板子和螢幕之間，電池線繞到左邊的插座。", PINK),
                                   ("記憶卡", "卡插到底還會露出一小段，剛好穿過殼側面的口，用指甲就能拔。", AMBER)]):
    x = 0.70 + i * 4.05
    rect(s, x, 6.15, 3.85, 0.75)
    square(s, x + 0.15, 6.27, col, 0.18)
    tbox(s, x + 0.42, 6.19, 1.3, 0.26, hd, 12, True, INK)
    tbox(s, x + 1.55, 6.19, 2.20, 0.70, bd, 9, False, INK2, space_after=0)
ud.finish(prs, ud.FOOTER_SIMPLE, page_x=11.60, footer_x=0.7)
prs.save("esp32s3-camera-v1.4-簡易版.pptx"); print("simple deck:", len(prs.slides), "slides")
